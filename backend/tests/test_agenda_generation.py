import copy
import json
from unittest.mock import Mock

import pytest
from pptx import Presentation
from pptx.util import Inches

from app.api.v1.presentations import (
    preserve_visual_context,
    shift_deck_feedback_references,
)
from app.services import agenda_generator as generator
from app.services.agenda_layouts import (
    AGENDA_LAYOUTS,
    PREVIEW_ITEMS,
    agenda_catalog,
    agenda_scene,
)


@pytest.mark.parametrize("layout", list(AGENDA_LAYOUTS))
@pytest.mark.parametrize("count", [4, 5, 6])
@pytest.mark.parametrize("dimensions", [(13.333333, 7.5), (10, 5.625), (10, 7.5)])
def test_editable_layout_content_and_bounds_survive_save(tmp_path, layout, count, dimensions):
    deck = Presentation()
    deck.slide_width, deck.slide_height = map(Inches, dimensions)
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    # A content placeholder must not survive under the generated design.
    slide.shapes.title.text = "INHERITED PLACEHOLDER"
    items = PREVIEW_ITEMS[:count]
    generator.draw_agenda(slide, items, layout, "ocean", deck.slide_width, deck.slide_height)
    path = tmp_path / "agenda.pptx"
    deck.save(path)
    loaded = Presentation(path)
    text = " ".join(s.text.replace("\n", " ") for s in loaded.slides[0].shapes if s.has_text_frame)
    assert "INHERITED PLACEHOLDER" not in text
    for item in items:
        assert item["title"] in text
        assert item["description"] in text
    for shape in loaded.slides[0].shapes:
        assert shape.left >= 0 and shape.top >= 0
        assert shape.left + shape.width <= loaded.slide_width + 2
        assert shape.top + shape.height <= loaded.slide_height + 2
        if shape.has_text_frame and shape.text:
            assert shape.text_frame.margin_top == 0
            for paragraph in shape.text_frame.paragraphs:
                assert paragraph.font.size is not None
                assert paragraph.line_spacing.pt < paragraph.font.size.pt * 1.21


def test_preview_uses_the_writer_layouts_and_full_palettes():
    catalog = agenda_catalog()
    assert len(catalog["templates"]) == 5
    assert all(
        len({palette[f"accent{i}"] for i in range(6)}) == 6
        for palette in catalog["themes"].values()
    )
    for template in catalog["templates"]:
        assert template["nodes"] == agenda_scene(PREVIEW_ITEMS, template["id"])


def test_agenda_refresh_preserves_existing_visual_context():
    refreshed = preserve_visual_context(
        {"num": 3, "text": [{"t": "Updated text"}], "img_context": []},
        {"num": 2, "img_context": "Two columns with a diagram on the right."},
    )
    assert refreshed["img_context"] == "Two columns with a diagram on the right."


def test_agenda_insertion_shifts_slide_numbers_inside_deck_feedback():
    shifted = shift_deck_feedback_references(
        {
            "score_rationale": "Slides 1 and 2 establish context before slide 4.",
            "deck_top_priorities": ["[HIGH] Slides 2-4 need a clearer transition."],
            "examiner_red_flags": ["The claim on slide 5 is unsupported."],
            "cohesion_feedback": "Slide 1 works.",
        },
        insert_after_slide=1,
    )

    assert shifted["score_rationale"] == "Slides 1 and 3 establish context before slide 5."
    assert shifted["deck_top_priorities"] == ["[HIGH] Slides 3-5 need a clearer transition."]
    assert shifted["examiner_red_flags"] == ["The claim on slide 6 is unsupported."]
    assert shifted["cohesion_feedback"] == "Slide 1 works."


@pytest.mark.parametrize("layout", list(AGENDA_LAYOUTS))
def test_long_sections_fit_without_overlapping_other_text(layout):
    items = [
        {
            "title": f"Research findings and implications {i}",
            "description": (
                "The important research findings that inform our next development steps."
            ),
        }
        for i in range(6)
    ]
    text_nodes = [node for node in agenda_scene(items, layout) if node["kind"] == "text"]
    for index, node in enumerate(text_nodes):
        assert len(node["lines"]) * node["size"] * 1.2 <= node["h"]
        for other in text_nodes[index + 1 :]:
            overlaps_x = min(node["x"] + node["w"], other["x"] + other["w"]) > max(
                node["x"], other["x"]
            )
            overlaps_y = min(node["y"] + node["h"], other["y"] + other["h"]) > max(
                node["y"], other["y"]
            )
            assert not (overlaps_x and overlaps_y)


def model_response(items):
    return Mock(
        json=lambda: {
            "choices": [
                {
                    "message": {
                        "content": json.dumps(
                            {
                                "items": items,
                                "insert_after_slide": 1,
                            }
                        )
                    }
                }
            ]
        }
    )


def test_missing_description_retries_instead_of_returning_title_only(monkeypatch):
    incomplete = copy.deepcopy(PREVIEW_ITEMS)
    incomplete[0].pop("description")
    post = Mock(side_effect=[model_response(incomplete), model_response(PREVIEW_ITEMS)])
    monkeypatch.setattr(generator.requests, "post", post)
    monkeypatch.setattr(generator, "openrouter_headers", lambda: {})
    monkeypatch.setattr(generator.time, "sleep", lambda _: None)
    items, position = generator._items_from_model([{"num": 1}])
    assert items == PREVIEW_ITEMS and position == 1
    assert post.call_count == 2


def test_invalid_model_output_fails_after_three_attempts_without_writing(monkeypatch, tmp_path):
    too_long = copy.deepcopy(PREVIEW_ITEMS)
    too_long[0]["title"] = "This title is far too long to use in a compact agenda section"
    post = Mock(return_value=model_response(too_long))
    monkeypatch.setattr(generator.requests, "post", post)
    monkeypatch.setattr(generator, "openrouter_headers", lambda: {})
    monkeypatch.setattr(generator.time, "sleep", lambda _: None)
    destination = tmp_path / "output.pptx"
    with pytest.raises(generator.AgendaGenerationError):
        generator.generate_agenda_deck(
            tmp_path / "unused.pptx", destination, [{"num": 1}], "paired-sections", "violet"
        )
    assert post.call_count == 3
    assert not destination.exists()


def test_agenda_inserted_at_model_position_preserves_other_slides(monkeypatch, tmp_path):
    source = tmp_path / "source.pptx"
    output = tmp_path / "output.pptx"
    deck = Presentation()
    for title in ("Cover", "Research", "Findings"):
        deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = title
    deck.save(source)
    monkeypatch.setattr(generator, "_items_from_model", lambda _: (PREVIEW_ITEMS, 1))
    titles, inserted_after = generator.generate_agenda_deck(
        source,
        output,
        [{"num": n} for n in range(1, 4)],
        "statement-grid",
        "forest",
    )
    result = Presentation(output)
    assert len(Presentation(source).slides) == 3
    assert len(result.slides) == 4
    assert inserted_after == 1
    assert titles == [item["title"] for item in PREVIEW_ITEMS]
    assert [result.slides[i].shapes.title.text for i in (0, 2, 3)] == [
        "Cover",
        "Research",
        "Findings",
    ]
    assert result.slides[1]._element.get("showMasterSp") == "0"
