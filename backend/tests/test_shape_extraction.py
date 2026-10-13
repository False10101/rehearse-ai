import logging

import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

from app.services.slide_analyzer import extract_slides
from app.services.slide_design import shape_evidence, shape_geometry


def make_slide():
    deck = Presentation()
    return deck, deck.slides.add_slide(deck.slide_layouts[6])


def set_group_transform(group, off, ext, child_off, child_ext, **attributes):
    node = group._element.find("./" + qn("p:grpSpPr") + "/" + qn("a:xfrm"))
    for name, values, keys in (
        ("off", off, ("x", "y")), ("ext", ext, ("cx", "cy")),
        ("chOff", child_off, ("x", "y")), ("chExt", child_ext, ("cx", "cy")),
    ):
        child = node.find(qn("a:" + name))
        for key, value in zip(keys, values):
            child.set(key, str(value))
    for key, value in attributes.items():
        node.set(key, str(value))


def test_unmapped_line_and_arrowheads_survive_full_extraction(tmp_path, caplog):
    deck, slide = make_slide()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 100, 200, 300, 0)
    shape._element.spPr.prstGeom.set("prst", "line")
    line = shape._element.spPr.get_or_add_ln()
    arrow = OxmlElement("a:tailEnd")
    arrow.set("type", "triangle")
    line.append(arrow)
    path = tmp_path / "line.pptx"
    deck.save(path)
    with caplog.at_level(logging.WARNING):
        evidence = extract_slides(path)[0]["shape_evidence"][0]
    assert evidence["raw_geometry"] == evidence["geometry"] == "line"
    assert evidence["line"]["slide_endpoints"] == [[100, 200], [400, 200]]
    assert evidence["line"]["tail_end"]["type"] == "triangle"
    assert "geometry=line" in caplog.text


def test_normal_rectangle_and_unknown_preset_are_retained():
    _, slide = make_slide()
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 100, 200, 300, 400)
    assert shape_evidence(shape, slide, 1)["geometry"].startswith("RECTANGLE")
    shape._element.spPr.prstGeom.set("prst", "futureGeometry")
    assert shape_evidence(shape, slide, 1)["geometry"] == "futureGeometry"


def test_connector_flips_preserve_arrow_endpoint_order_and_links():
    _, slide = make_slide()
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 400, 200, 100, 500)
    assert shape_evidence(connector, slide, 1)["line"]["slide_endpoints"] == [
        [400, 200], [100, 500]
    ]
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 100, 500, 200, 200)
    connector.end_connect(box, 0)
    data = shape_evidence(connector, slide, 1)
    assert data["element_type"] == "cxnSp"
    assert data["flip_horizontal"]
    assert data["line"]["explicit_connections"][0]["id"] == str(box.shape_id)
    assert len(data["line"]["slide_endpoints"]) == 2


def test_nested_groups_project_positions_without_losing_local_bounds():
    _, slide = make_slide()
    outer = slide.shapes.add_group_shape()
    inner = outer.shapes.add_group_shape()
    shape = inner.shapes.add_shape(MSO_SHAPE.RECTANGLE, 10, 20, 30, 40)
    set_group_transform(inner, (100, 200), (200, 300), (0, 0), (100, 100))
    set_group_transform(outer, (1000, 2000), (400, 600), (0, 0), (200, 300))
    data = shape_geometry(shape)
    assert [data[k] for k in ("left", "top", "width", "height")] == [1240, 2520, 120, 240]
    assert data["local_bounds"] == {"left": 10, "top": 20, "width": 30, "height": 40}
    assert len(data["parent_group_ids"]) == 2


def test_shape_rotation_and_group_flip():
    _, slide = make_slide()
    group = slide.shapes.add_group_shape()
    shape = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, 10, 20, 30, 40)
    shape.rotation = 90
    set_group_transform(group, (0, 0), (100, 100), (0, 0), (100, 100), flipH="1")
    data = shape_geometry(shape)
    assert [data[k] for k in ("left", "top", "width", "height")] == [55, 25, 40, 30]


def test_group_rotation_projects_line_endpoints():
    _, slide = make_slide()
    group = slide.shapes.add_group_shape()
    shape = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 50, 100, 0)
    shape._element.spPr.prstGeom.set("prst", "line")
    set_group_transform(group, (0, 0), (100, 100), (0, 0), (100, 100), rot="5400000")
    assert shape_evidence(shape, slide, 1)["line"]["slide_endpoints"] == [[50, 0], [50, 100]]


def test_zero_height_line_group_is_valid():
    _, slide = make_slide()
    group = slide.shapes.add_group_shape()
    shape = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, 10, 20, 30, 0)
    shape._element.spPr.prstGeom.set("prst", "line")
    set_group_transform(group, (100, 200), (60, 0), (10, 20), (30, 0))
    assert shape_evidence(shape, slide, 1)["line"]["slide_endpoints"] == [[100, 200], [160, 200]]


def test_invalid_group_transform_is_not_silently_accepted():
    _, slide = make_slide()
    group = slide.shapes.add_group_shape()
    shape = group.shapes.add_shape(MSO_SHAPE.RECTANGLE, 10, 20, 30, 40)
    set_group_transform(group, (0, 0), (100, 100), (0, 0), (0, 100))
    with pytest.raises(ValueError, match="non-projectable"):
        shape_geometry(shape)


def test_custom_shape_paths_are_retained():
    _, slide = make_slide()
    builder = slide.shapes.build_freeform(0, 0)
    builder.add_line_segments([(100, 0), (100, 100), (0, 0)])
    shape = builder.convert_to_shape()
    evidence = shape_evidence(shape, slide, 1)
    assert evidence["geometry"] == "custom"
    assert "pathLst" in evidence["custom_geometry_xml"]
