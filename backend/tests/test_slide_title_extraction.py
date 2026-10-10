from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

from app.services.slide_analyzer import extract_slides


def test_extract_slides_uses_a_powerpoint_title_placeholder(tmp_path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    slide.shapes.title.text = "Research Findings"
    slide.placeholders[1].text = "A short summary"
    filepath = tmp_path / "title-placeholder.pptx"
    deck.save(filepath)

    slides = extract_slides(filepath)

    assert slides[0]["title"] == "Research Findings"


def test_extract_slides_recurses_into_grouped_text_and_preserves_structure(tmp_path: Path) -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    text_box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(2))
    text_box.text_frame.paragraphs[0].text = "AI Matching"
    second = text_box.text_frame.add_paragraph()
    second.text = "Compares skills with role requirements."
    second.level = 1
    filepath = tmp_path / "grouped-text.pptx"
    deck.save(filepath)

    extracted = extract_slides(filepath)[0]

    assert "AI Matching" in "".join(run["t"] for run in extracted["text"])
    assert "Compares skills" in "".join(run["t"] for run in extracted["text"])
    assert extracted["imgs"] == 0
    assert extracted["paragraphs"][1]["level"] == 1
    assert extracted["text_regions"][0]["width"] > 0
