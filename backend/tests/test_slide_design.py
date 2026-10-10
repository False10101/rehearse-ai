from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from app.services.slide_analyzer import deterministic_table_tips, extract_slides
from app.services.slide_design import design_findings, overlap_candidates


def test_paragraph_defaults_and_local_panel_contrast(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    panel = slide.shapes.add_shape(1, Inches(1), Inches(1), Inches(4), Inches(3))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(0, 0, 0)
    for x in (1.2, 5.2):
        shape = slide.shapes.add_textbox(Inches(x), Inches(1.5), Inches(2), Inches(0.5))
        paragraph = shape.text_frame.paragraphs[0]
        paragraph.font.size = Pt(24)
        paragraph.font.bold = True
        paragraph.font.name = "Arial"
        paragraph.font.color.rgb = RGBColor(255, 255, 255)
        paragraph.text = "Completion metric"
    path = tmp_path / "contrast.pptx"
    deck.save(path)
    extracted = extract_slides(path)[0]
    text = [r for r in extracted["text"] if r["t"] == "Completion metric"]
    assert all(
        r["sz"] == 24 and r["bold"] and r["font_family"] == "Arial" and r["color"] == "#FFFFFF"
        for r in text
    )
    findings = design_findings(extracted)
    assert len(findings) == 1  # white on black is not a fault; white on white is
    assert "#FFFFFF on #FFFFFF" in findings[0]


def test_ambiguous_prices_do_not_invent_discount_and_limits_conflict():
    tips = deterministic_table_tips({"tables": [[["$12 or $140?", "Unlimited (max 8)"]]]})
    assert not any("regular and discounted" in tip for tip in tips)
    assert any("relationship" in tip for tip in tips)
    assert any("numeric limit" in tip for tip in tips)


def test_unknown_background_does_not_claim_contrast():
    slide = {
        "text_regions": [
            {
                "region": 1,
                "left": 0,
                "top": 0,
                "width": 20,
                "height": 20,
                "paragraphs": [{"text": "hello"}],
                "runs": [{"t": "hello", "color": "#FFFFFF"}],
            }
        ]
    }
    assert design_findings(slide) == []


def test_low_contrast_beyond_nearly_invisible_text(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    for y, color in ((1, (170, 170, 170)), (2, (0, 0, 0))):
        box = slide.shapes.add_textbox(Inches(1), Inches(y), Inches(5), Inches(1))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = "Status detail" if y == 1 else "Readable detail"
        paragraph.font.size = Pt(24)
        paragraph.font.color.rgb = RGBColor(*color)
    path = tmp_path / "weak-contrast.pptx"
    deck.save(path)
    findings = design_findings(extract_slides(path)[0])
    assert len(findings) == 1
    assert findings[0].startswith("[MEDIUM]")
    assert "Status detail" in findings[0]


def test_wording_and_typography_are_independent_prompt_checks():
    from app.services.slide_analyzer import build_analysis_prompt
    prompt = build_analysis_prompt([{"num": 1, "title": "Example"}])
    assert "circular definitions, tautologies" in prompt
    assert "Scope/overload feedback does not substitute for wording feedback" in prompt
    assert "Reducing scope alone does not address typography" in prompt


def test_foreground_overlap_is_candidate_not_automatic_fault(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    panel = slide.shapes.add_shape(1, 0, 0, Inches(5), Inches(2))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(255, 255, 255)
    textbox = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(4), Inches(1))
    textbox.text = "Approval required before launch"
    cover = slide.shapes.add_shape(1, Inches(2), Inches(0.5), Inches(2), Inches(1))
    cover.fill.solid()
    cover.fill.fore_color.rgb = RGBColor(220, 0, 0)
    path = tmp_path / "overlap.pptx"
    deck.save(path)
    extracted = extract_slides(path)[0]
    candidates = overlap_candidates(extracted)
    assert len(candidates) == 1
    assert candidates[0]["text_region"] == 2
    assert candidates[0]["foreground_shape"] == 3
    assert candidates[0]["box_overlap_fraction"] == 0.5
    assert candidates[0]["status"] == "candidate_requires_visual_confirmation"
    assert not any("obscur" in item for item in design_findings(extracted))


def test_overlap_crops_use_one_existing_vision_request(monkeypatch, tmp_path):
    from PIL import Image

    from app.services import slide_analyzer

    image = tmp_path / "render.png"
    Image.new("RGB", (800, 450), "white").save(image)
    shape = {
        "z_order": 1,
        "left": 10,
        "top": 10,
        "width": 100,
        "height": 30,
        "coordinate_scope": "slide",
        "rotation": 0,
        "fill_color": None,
    }
    evidence = {
        "slide_width": 200,
        "slide_height": 100,
        "text_regions": [
            {
                "region": 1,
                "left": 10,
                "top": 10,
                "width": 100,
                "height": 30,
                "paragraphs": [{"text": "Approval required"}],
            }
        ],
        "shape_evidence": [
            shape,
            {**shape, "z_order": 2, "left": 30, "width": 40, "fill_color": "#FF0000"},
        ],
    }
    calls = []

    class Response:
        def raise_for_status(self):
            pass

        def json(self):
            import json

            return {
                "choices": [
                    {
                        "message": {
                            "content": json.dumps(
                                {
                                    "description": "A text region is visible.",
                                    "checks": [
                                        {
                                            "check_id": "hierarchy",
                                            "status": "not_issue",
                                            "reason": "One text region has clear prominence.",
                                        },
                                        {
                                            "check_id": "overlap-1-2",
                                            "status": "uncertain",
                                            "reason": "Coverage cannot be established.",
                                        }
                                    ],
                                }
                            )
                        }
                    }
                ]
            }

    def post(*args, **kwargs):
        calls.append(kwargs["json"])
        return Response()

    monkeypatch.setattr(slide_analyzer, "openrouter_headers", lambda: {})
    monkeypatch.setattr(slide_analyzer.requests, "post", post)
    slide_analyzer.caption_slide_image(1, image, "Approval required", evidence)
    assert len(calls) == 1
    content = calls[0]["messages"][0]["content"]
    assert len([item for item in content if item["type"] == "image_url"]) == 2
    assert "candidate_requires_visual_confirmation" in content[1]["text"]


def test_theme_color_and_layout_placeholder_inheritance(tmp_path):
    deck = Presentation()
    layout = deck.slide_layouts[0]
    title_default = layout.placeholders[0].text_frame.paragraphs[0].font
    title_default.size = Pt(37)
    color = OxmlElement("a:solidFill")
    scheme = OxmlElement("a:schemeClr")
    scheme.set("val", "dk1")
    color.append(scheme)
    title_default._rPr.append(color)
    slide = deck.slides.add_slide(layout)
    slide.shapes.title.text = "Inherited heading"
    path = tmp_path / "theme.pptx"
    deck.save(path)
    run = next(r for r in extract_slides(path)[0]["text"] if r["t"] == "Inherited heading")
    assert run["sz"] == 37
    assert run["color"] == "#000000"


def test_table_font_defaults_are_retained(tmp_path):
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    table = slide.shapes.add_table(1, 1, 0, 0, Inches(2), Inches(1)).table
    cell = table.cell(0, 0)
    cell.text = "Plan limit"
    cell.text_frame.paragraphs[0].font.size = Pt(19)
    path = tmp_path / "table.pptx"
    deck.save(path)
    assert extract_slides(path)[0]["table_regions"][0]["cell_styles"][0]["runs"][0]["sz"] == 19


def test_panel_comparison_reports_measurements_not_faults():
    from app.services.slide_design import alignment_evidence

    shapes = [
        {
            "z_order": i + 1,
            "left": x,
            "top": y,
            "width": w,
            "height": 100,
            "fill_color": "#AAAAAA",
            "geometry": "RECTANGLE (1)",
            "coordinate_scope": "slide",
            "rotation": 0,
        }
        for i, (x, y, w) in enumerate([(0, 0, 100), (120, 15, 100), (260, 0, 140)])
    ]
    evidence = alignment_evidence({"slide_width": 500, "shape_evidence": shapes})
    assert len(evidence) == 1
    assert evidence[0]["width_spread_fraction"] == 0.08
    assert evidence[0]["gap_fractions"] == [0.04, 0.08]
    assert evidence[0]["status"] == "comparison_requires_visual_review"
    assert alignment_evidence({"shape_evidence": shapes[:2]}) == []


def test_style_summary_and_small_text_are_review_evidence():
    from app.services.slide_design import deck_style_evidence, typography_evidence

    def slide(num, font, background):
        return {
            "num": num,
            "slide_height": 100,
            "bg_color": background,
            "text_regions": [
                {
                    "region": 1,
                    "top": 35,
                    "paragraphs": [
                        {
                            "text": "This body passage contains enough words "
                            "to compare its styling with other content slides."
                        }
                    ],
                    "runs": [{"t": "body", "sz": 14, "color": "#000000", "font_family": font}],
                }
            ],
        }

    slides = [slide(i, "Arial", "#FFFFFF") for i in range(1, 4)]
    slides.append(slide(4, "Georgia", "#FF00FF"))
    summary = deck_style_evidence(slides)
    assert summary["repeated_content_styles"]["body_fonts"]["supporting_slides"] == [1, 2, 3]
    assert summary["profiles"][3]["body_fonts"] == ["Georgia"]
    assert typography_evidence(slides[0])[0]["readability_review_candidate"]
    assert "not faults" in summary["rule"]
    slides[3]["generated_by"] = "agenda"
    assert len(deck_style_evidence(slides)["profiles"]) == 3


def test_pre_generation_evidence_does_not_mutate_input():
    from app.services.slide_analyzer import build_analysis_prompt

    slide = {"num": 1, "title": "Example"}
    prompt = build_analysis_prompt([slide])
    assert "objective_visual_findings" in prompt
    assert "DECK_STYLE_EVIDENCE" in prompt
    assert "typography_review" in prompt
    assert slide == {"num": 1, "title": "Example"}


def test_contrast_fallback_does_not_repeat_equivalent_model_advice(monkeypatch):
    from app.services import slide_analyzer

    monkeypatch.setattr(
        slide_analyzer,
        "design_findings",
        lambda _: [
            "[HIGH] Increase foreground/background contrast for “Hidden metric” "
            "(#FFFFFF on #FFFFFF)."
        ],
    )
    result = {
        "per_slide": [
            {
                "num": 1,
                "table_tips": [],
                "improvement_suggestions": [
                    "[HIGH] Give “Hidden metric” a darker text color.",
                    "[MEDIUM] Explain what the metric measures.",
                ],
                "visual_hierarchy_issues": "The metric needs darker text.",
            }
        ]
    }
    slide_analyzer.apply_deterministic_findings(result, [{"num": 1}])
    assert len(result["per_slide"][0]["improvement_suggestions"]) == 2
    assert result["per_slide"][0]["visual_hierarchy_issues"] == "The metric needs darker text."
    result["per_slide"][0]["improvement_suggestions"] = ["[MEDIUM] Explain the metric."]
    slide_analyzer.apply_deterministic_findings(result, [{"num": 1}])
    assert len(result["per_slide"][0]["improvement_suggestions"]) == 2
