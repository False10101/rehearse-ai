"""Conservative PowerPoint style and geometry evidence, not pixel-perfect rendering."""

from collections import Counter

from lxml import etree
from pptx.oxml.ns import qn


def _theme(slide):
    for relation in slide.slide_layout.slide_master.part.rels.values():
        if relation.reltype.endswith("/theme"):
            return etree.fromstring(relation.target_part.blob)
    return None


def _color(node, theme):
    if node is None:
        return None
    # Transformed/transparent colors need renderer evaluation; do not guess.
    if len(node):
        return None
    if node.tag == qn("a:srgbClr"):
        return "#" + node.get("val")
    if node.tag == qn("a:schemeClr") and theme is not None:
        name = {"tx1": "dk1", "tx2": "dk2", "bg1": "lt1", "bg2": "lt2"}.get(
            node.get("val"), node.get("val")
        )
        scheme = theme.find(".//" + qn("a:clrScheme"))
        entry = scheme.find(qn("a:" + name)) if scheme is not None else None
        if entry is not None and len(entry):
            value = (
                entry[0].get("val") if entry[0].tag == qn("a:srgbClr") else entry[0].get("lastClr")
            )
            return "#" + value if value else None
    return None


def resolved_style(run, paragraph, shape, slide):
    """Resolve each property separately from explicit run through inherited defaults."""
    theme = _theme(slide)
    nodes = []
    if run._r.rPr is not None:
        nodes.append(run._r.rPr)
    level = paragraph.level + 1
    sources = [shape]
    if getattr(shape, "is_placeholder", False):
        idx = shape.placeholder_format.idx
        layout_shape = next(
            (p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == idx), None
        )
        if layout_shape is not None:
            sources.append(layout_shape)
            kind = layout_shape.placeholder_format.type
            master_shape = next(
                (
                    p
                    for p in slide.slide_layout.slide_master.placeholders
                    if p.placeholder_format.type == kind
                ),
                None,
            )
            if master_shape is not None:
                sources.append(master_shape)
    for source in sources:
        paragraphs = source.text_frame.paragraphs
        p = paragraph if source is shape else paragraphs[0]
        if p._p.pPr is not None:
            default = p._p.pPr.find(qn("a:defRPr"))
            if default is not None:
                nodes.append(default)
        for tag in (f"a:lvl{level}pPr", "a:defPPr"):
            nodes.extend(
                source.text_frame._txBody.findall(
                    ".//" + qn("a:lstStyle") + "/" + qn(tag) + "/" + qn("a:defRPr")
                )
            )
    kind = "otherStyle"
    if getattr(shape, "is_placeholder", False):
        name = str(shape.placeholder_format.type)
        kind = "titleStyle" if "TITLE" in name else "bodyStyle"
    master = slide.slide_layout.slide_master._element
    nodes.extend(
        master.findall(
            ".//" + qn("p:" + kind) + "/" + qn(f"a:lvl{level}pPr") + "/" + qn("a:defRPr")
        )
    )
    nodes.extend(
        slide.part.package.presentation_part._element.findall(
            ".//" + qn("p:defaultTextStyle") + "/" + qn(f"a:lvl{level}pPr") + "/" + qn("a:defRPr")
        )
    )
    result = {"sz": None, "bold": None, "color": None, "font_family": None}
    for node in nodes:
        if result["sz"] is None and node.get("sz"):
            result["sz"] = int(node.get("sz")) / 100
        if result["bold"] is None and node.get("b") is not None:
            result["bold"] = node.get("b") in ("1", "true")
        fill = node.find(qn("a:solidFill"))
        if result["color"] is None and fill is not None and len(fill):
            result["color"] = _color(fill[0], theme)
            # An unresolved explicit fill must not inherit a different foreground.
            if result["color"] is None:
                result["color"] = "unknown"
        latin = node.find(qn("a:latin"))
        if result["font_family"] is None and latin is not None:
            result["font_family"] = latin.get("typeface")
    family = result["font_family"]
    if family in ("+mj-lt", "+mn-lt") and theme is not None:
        tag = "majorFont" if family == "+mj-lt" else "minorFont"
        latin = theme.find(".//" + qn("a:" + tag) + "/" + qn("a:latin"))
        result["font_family"] = latin.get("typeface") if latin is not None else None
    if result["color"] == "unknown":
        result["color"] = None
    return result


def shape_evidence(shape, slide, index):
    fill = shape._element.find("./" + qn("p:spPr") + "/" + qn("a:solidFill"))
    color = _color(fill[0], _theme(slide)) if fill is not None and len(fill) else None
    return {
        "z_order": index,
        "left": shape.left,
        "top": shape.top,
        "coordinate_scope": "group" if shape._element.getparent().tag == qn("p:grpSp") else "slide",
        "width": shape.width,
        "height": shape.height,
        "fill_color": color,
        "geometry": str(getattr(shape, "auto_shape_type", "unknown"))
        if shape.shape_type == 1
        else "other",
        "rotation": shape.rotation,
        "text": shape.text if shape.has_text_frame else "",
    }


def overlap_candidates(slide):
    """Find later solid shapes intersecting earlier text, not confirmed hidden glyphs."""
    candidates = []
    for region in slide.get("text_regions", []):
        text = " ".join(p["text"] for p in region["paragraphs"]).strip()
        if not text or region["width"] <= 0 or region["height"] <= 0:
            continue
        own = next(
            (s for s in slide.get("shape_evidence", []) if s["z_order"] == region["region"]), None
        )
        if not own or own["coordinate_scope"] != "slide" or own["rotation"]:
            continue
        for shape in slide.get("shape_evidence", []):
            if (
                shape["z_order"] <= region["region"]
                or not shape["fill_color"]
                or shape["coordinate_scope"] != "slide"
                or shape["rotation"]
            ):
                continue
            left = max(region["left"], shape["left"])
            top = max(region["top"], shape["top"])
            right = min(region["left"] + region["width"], shape["left"] + shape["width"])
            bottom = min(region["top"] + region["height"], shape["top"] + shape["height"])
            if right <= left or bottom <= top:
                continue
            fraction = (right - left) * (bottom - top) / (region["width"] * region["height"])
            candidates.append(
                {
                    "text_region": region["region"],
                    "foreground_shape": shape["z_order"],
                    "source_text": text,
                    "box_overlap_fraction": round(fraction, 3),
                    "crop_bounds": [
                        min(region["left"], shape["left"]),
                        min(region["top"], shape["top"]),
                        max(region["left"] + region["width"], shape["left"] + shape["width"]),
                        max(region["top"] + region["height"], shape["top"] + shape["height"]),
                    ],
                    "status": "candidate_requires_visual_confirmation",
                }
            )
    return sorted(candidates, key=lambda item: item["box_overlap_fraction"], reverse=True)


def typography_evidence(slide):
    """Measured text styles; sizes alone are not proof of illegibility."""
    regions = []
    for region in slide.get("text_regions", []):
        runs = [r for r in region.get("runs", []) if r.get("t", "").strip()]
        if not runs:
            continue
        sizes = sorted({r["sz"] for r in runs if r.get("sz")})
        text = " ".join(p["text"] for p in region["paragraphs"]).strip()
        regions.append(
            {
                "region": region["region"],
                "source_text": text[:200],
                "word_count": len(text.split()),
                "font_sizes_pt": sizes,
                "font_families": sorted({r["font_family"] for r in runs if r.get("font_family")}),
                "colors": sorted({r["color"] for r in runs if r.get("color")}),
                "readability_review_candidate": bool(sizes and min(sizes) < 18),
                "top_fraction": round(region["top"] / slide["slide_height"], 3)
                if slide.get("slide_height")
                else None,
            }
        )
    return regions


def alignment_evidence(slide):
    """Compare likely repeated horizontal panels, without judging intentional asymmetry."""
    panels = [
        s
        for s in slide.get("shape_evidence", [])
        if s.get("fill_color")
        and s.get("geometry", "").startswith("RECTANGLE")
        and s.get("coordinate_scope") == "slide"
        and not s.get("rotation")
        and s["width"] > 0
        and s["height"] > 0
    ]
    groups = []
    seen = set()
    for anchor in panels:
        row = sorted(
            [
                p
                for p in panels
                if p["fill_color"] == anchor["fill_color"]
                and abs(p["height"] - anchor["height"]) <= anchor["height"] * 0.15
                and abs(p["top"] - anchor["top"]) <= anchor["height"] * 0.5
            ],
            key=lambda p: p["left"],
        )
        ids = tuple(p["z_order"] for p in row)
        if len(row) < 3 or ids in seen:
            continue
        if any(a["left"] + a["width"] > b["left"] for a, b in zip(row, row[1:])):
            continue  # overlapping/nested rectangles are not a repeated-card row
        seen.add(ids)
        scale = slide.get("slide_width") or 1
        groups.append(
            {
                "shape_ids": list(ids),
                "status": "comparison_requires_visual_review",
                "tops": [p["top"] for p in row],
                "widths": [p["width"] for p in row],
                "heights": [p["height"] for p in row],
                "gap_fractions": [
                    round((b["left"] - a["left"] - a["width"]) / scale, 4)
                    for a, b in zip(row, row[1:])
                ],
                "top_spread_fraction": round(
                    (max(p["top"] for p in row) - min(p["top"] for p in row)) / scale, 4
                ),
                "width_spread_fraction": round(
                    (max(p["width"] for p in row) - min(p["width"] for p in row)) / scale, 4
                ),
            }
        )
    return groups


def deck_style_evidence(slides):
    """Expose comparable style patterns; do not declare every outlier defective."""
    profiles = []
    for slide in slides:
        if slide.get("generated_by") == "agenda":
            continue
        regions = typography_evidence(slide)
        heading = [
            r
            for r in regions
            if r["top_fraction"] is not None and r["top_fraction"] < 0.25 and r["word_count"] <= 12
        ]
        body = [r for r in regions if r not in heading]
        content_candidate = sum(r["word_count"] for r in body) >= 12 or bool(slide.get("tables"))
        profiles.append(
            {
                "num": slide["num"],
                "content_comparison_candidate": content_candidate,
                "background": slide.get("bg_color"),
                "heading_fonts": sorted({f for r in heading for f in r["font_families"]}),
                "body_fonts": sorted({f for r in body for f in r["font_families"]}),
                "heading_colors": sorted({c for r in heading for c in r["colors"]}),
                "body_colors": sorted({c for r in body for c in r["colors"]}),
            }
        )
    comparable = [p for p in profiles if p["content_comparison_candidate"]]
    common = {}
    for field in ("background", "heading_fonts", "body_fonts", "heading_colors"):
        counts = Counter(
            tuple(p[field]) if isinstance(p[field], list) else p[field]
            for p in comparable
            if p[field]
        )
        if counts:
            value, count = counts.most_common(1)[0]
            if count >= 3:
                common[field] = {
                    "value": value,
                    "supporting_slides": [
                        p["num"]
                        for p in comparable
                        if (tuple(p[field]) if isinstance(p[field], list) else p[field]) == value
                    ],
                }
    return {
        "profiles": profiles,
        "repeated_content_styles": common,
        "rule": "Role assignments are heuristic. Differences are review candidates, not faults; "
        "allow intentional section, title, closing and accessible emphasis variations.",
    }


def design_findings(slide):
    """Emit definite overflow and low contrast on a known solid local background."""
    findings = []
    width, height = slide.get("slide_width"), slide.get("slide_height")
    shapes = slide.get("shape_evidence", [])
    for region in slide.get("text_regions", []):
        text = " ".join(p["text"] for p in region["paragraphs"]).strip()
        if not text:
            continue
        own_shape = next((s for s in shapes if s["z_order"] == region["region"]), None)
        if own_shape and (own_shape["coordinate_scope"] != "slide" or own_shape["rotation"]):
            continue  # local group coordinates and rotated boxes need rendering confirmation
        label = text[:100]
        if (
            width
            and height
            and (
                region["left"] < 0
                or region["top"] < 0
                or region["left"] + region["width"] > width
                or region["top"] + region["height"] > height
            )
        ):
            findings.append(
                f"[HIGH] Move the text box containing “{label}” inside the slide "
                "boundaries and check the preview for cut-off content."
            )
        background = slide.get("bg_color")
        # Require fully containing, unrotated rectangles; intersecting/complex shapes are uncertain.
        for shape in shapes:
            if shape["z_order"] >= region["region"]:
                break
            intersects = (
                shape["left"] < region["left"] + region["width"]
                and shape["left"] + shape["width"] > region["left"]
                and shape["top"] < region["top"] + region["height"]
                and shape["top"] + shape["height"] > region["top"]
            )
            if not intersects:
                continue
            contains = (
                shape["left"] <= region["left"]
                and shape["top"] <= region["top"]
                and shape["left"] + shape["width"] >= region["left"] + region["width"]
                and shape["top"] + shape["height"] >= region["top"] + region["height"]
            )
            known_panel = (
                contains
                and shape["rotation"] == 0
                and shape["geometry"].startswith("RECTANGLE")
                and shape["coordinate_scope"] == "slide"
            )
            background = shape["fill_color"] if known_panel else None
        if own_shape and own_shape["fill_color"]:
            background = own_shape["fill_color"]
        if not background:
            continue
        for run in region.get("runs", []):
            color = run.get("color")
            if not color or not run["t"].strip():
                continue

            def luminance(rgb):
                channels = [int(rgb[i : i + 2], 16) / 255 for i in (1, 3, 5)]
                linear = [
                    c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4 for c in channels
                ]
                return sum(c * w for c, w in zip(linear, (0.2126, 0.7152, 0.0722), strict=True))

            a, b = sorted((luminance(color), luminance(background)))
            ratio = (b + 0.05) / (a + 0.05)
            # Below 3:1 fails even the large-text contrast benchmark. Do not
            # extend this conclusion to unknown, image or transparent backgrounds.
            if ratio < 3:
                excerpt = run["t"][:100]
                priority = "HIGH" if ratio < 1.5 else "MEDIUM"
                findings.append(
                    f"[{priority}] Increase foreground/background contrast for “{excerpt}” "
                    f"({color} on {background}); use a contrasting text or panel color "
                    "so this content can be read."
                )
    return list(dict.fromkeys(findings))
